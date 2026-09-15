"""AnalysisController —— 活动分析页的数据编排（阶段 4）。

从 MainWindow 迁出：热力图统计、活动报告树、活动内容视图、周期选择与报告导出。

页面部件由 :mod:`src.ui.views.analysis_view` 构建后经 :meth:`attach` 注入——
控制器只持有自己真正用到的部件，不反向读取 MainWindow 的私有属性。

分析页是懒构建的（首次切换到该页时才建），因此 ``attach()`` 之前所有部件均为
``None``，各槽方法都做了空值短路。
"""

from __future__ import annotations

from datetime import date

from PySide6.QtCore import QObject
from PySide6.QtWidgets import QFileDialog, QMessageBox

from ...services.task_service import TaskService

__all__ = ["AnalysisController"]


class AnalysisController(QObject):
    """活动分析页的槽函数与导出逻辑。

    生命周期：MainWindow 构造时创建 → 分析页首次构建时 :meth:`attach` 注入部件
    → 分区切换时 :meth:`set_partition` + :meth:`refresh`。
    """

    def __init__(
        self,
        task_service: TaskService,
        parent: QObject | None = None,
    ) -> None:
        super().__init__(parent)
        self._svc = task_service
        self._partition_id: str | None = None
        self._date_range: tuple = (None, None)

        # attach() 之前为 None（分析页懒构建）
        self._heatmap = None
        self._stats = None
        self._period = None
        self._tree = None
        self._content = None

    # ------------------------------------------------------------------
    # 装配 / 状态
    # ------------------------------------------------------------------

    def attach(
        self,
        *,
        heatmap,
        stats,
        period_selector,
        search,
        tree,
        content,
    ) -> None:
        """注入分析页部件（``analysis_view.build()`` 末尾调用）。"""
        self._heatmap = heatmap
        self._stats = stats
        self._period = period_selector
        self._search = search
        self._tree = tree
        self._content = content

    @property
    def is_attached(self) -> bool:
        """分析页是否已构建并注入部件。"""
        return self._tree is not None

    @property
    def date_range(self) -> tuple:
        """当前分析周期 ``(from, to)``；未选择时为 ``(None, None)``。"""
        return self._date_range

    def set_partition(self, partition_id: str | None) -> None:
        self._partition_id = partition_id or None

    # ------------------------------------------------------------------
    # 槽
    # ------------------------------------------------------------------

    def refresh(self, partition_id: str | None = None) -> None:
        """刷新热力图统计面板与活动报告树。"""
        if partition_id is not None:
            self._partition_id = partition_id or None

        if self._stats is not None and self._heatmap is not None:
            model = self._heatmap._model
            self._stats.refresh(
                total=model.total_count(),
                active_days=model.active_days(),
                longest_streak=model.longest_streak(),
                daily_avg=model.daily_average(),
            )

        if self._tree is not None:
            d_from, d_to = self._date_range
            self._tree.refresh(d_from, d_to, self._partition_id)

    def on_period_changed(self, d_from, d_to, label: str) -> None:
        """周期变化 → 高亮热力图区间 + 刷新报告树。"""
        self._date_range = (d_from, d_to)
        if self._heatmap is not None:
            if d_from is not None and d_to is not None:
                self._heatmap.highlight_range(d_from, d_to, label)
            else:
                self._heatmap.highlight_range(None, None, "")
        if self._tree is not None:
            self._tree.refresh(d_from, d_to, self._partition_id)

    def on_tag_selected(self, tag: str) -> None:
        """报告树选中标签 → 载入该标签的活动内容。"""
        if self._content is None:
            return
        if not tag:
            self._content.show_hint()
            return
        d_from, d_to = self._date_range
        tasks = self._tree.get_tasks_for_tag(tag) if self._tree is not None else []
        checked = self._tree.get_checked_tags() if self._tree is not None else []
        pos = checked.index(tag) + 1 if tag in checked else 0
        self._content.set_current_tag(tag, pos, len(checked))
        self._content.show_tag_activity(tag, tasks, d_from, d_to)

    def on_prev(self) -> None:
        if self._tree is not None:
            self._tree.select_prev()

    def on_next(self) -> None:
        if self._tree is not None:
            self._tree.select_next()

    def on_date_clicked(self, d: date) -> None:
        """热力图日期点击 → 切换为单日自定义周期。"""
        if self._period is not None:
            self._period.set_custom_range(d, d)

    def on_search_changed(self, text: str) -> None:
        """活动内容搜索。"""
        if self._content is not None:
            self._content.set_search_text(text)

    # ------------------------------------------------------------------
    # 导出
    # ------------------------------------------------------------------

    def export(self, fmt: str) -> None:
        """把全部勾选标签的活动内容导出为 md / xlsx / txt。"""
        d_from, d_to = self._date_range

        texts: list[str] = []
        checked_tags: list[str] = []
        if self._tree is not None and self._content is not None:
            checked_tags = self._tree.get_checked_tags()
            for tag in checked_tags:
                tasks = self._tree.get_tasks_for_tag(tag)
                self._content.show_tag_activity(tag, tasks, d_from, d_to)
                plain = self._content.get_plain_text()
                if plain:
                    texts.append(plain)
            # 还原用户当前查看的标签
            current_tag = self._tree.get_active_tag()
            if current_tag:
                tasks = self._tree.get_tasks_for_tag(current_tag)
                self._content.show_tag_activity(current_tag, tasks, d_from, d_to)

        text = "\n".join(texts)
        if not text:
            return

        def_name = self._build_export_filename(fmt, len(checked_tags))
        filters = {
            "md": "Markdown (*.md)",
            "xlsx": "Excel (*.xlsx)",
            "txt": "文本文件 (*.txt)",
        }
        parent = self._content
        filepath, _ = QFileDialog.getSaveFileName(
            parent, "导出报告", def_name, filters.get(fmt, "")
        )
        if not filepath:
            return

        if fmt == "xlsx":
            self._export_xlsx_file(filepath, text)
        else:
            with open(filepath, "w", encoding="utf-8") as fh:
                fh.write(text)

    def _build_export_filename(self, fmt: str, tag_count: int = 0) -> str:
        """默认文件名：``分区名_时间范围_N个标签.ext``。"""
        name_map = self._svc.get_partition_name_map()
        pname = name_map.get(self._partition_id or "", "默认分区")
        d_from, d_to = self._date_range
        date_str = ""
        if d_from and d_to:
            date_str = (
                f"{d_from.isoformat()}"
                if d_from == d_to
                else f"{d_from.isoformat()}~{d_to.isoformat()}"
            )
        tag_suffix = f"_{tag_count}个标签" if tag_count > 0 else ""
        return f"{pname}_{date_str}{tag_suffix}.{fmt}"

    def _export_xlsx_file(self, filepath: str, text: str = "") -> None:
        """导出 Excel（列：序号 / 任务 / 状态变更 / 进度变更 / 活动信息）。"""
        try:
            import openpyxl
            from openpyxl.styles import Font

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "活动报告"
            headers = ["序号", "任务", "状态变更", "进度变更", "活动信息"]
            for col, head in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=head)
                cell.font = Font(bold=True)
            for r, row_data in enumerate(self._parse_export_rows(text), 2):
                for c, val in enumerate(row_data, 1):
                    ws.cell(row=r, column=c, value=val)
            ws.column_dimensions["A"].width = 6
            ws.column_dimensions["B"].width = 30
            ws.column_dimensions["C"].width = 14
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 60
            wb.save(filepath)
        except ImportError:
            QMessageBox.warning(
                self._content, "错误", "需要安装 openpyxl 库才能导出 Excel"
            )

    def _parse_export_rows(self, text: str) -> list[tuple]:
        """把纯文本报告解析为 Excel 行。

        格式::

            #tag
            1. title [status, prog]:
                entry line 1
                entry line 2
        """
        rows = []
        current_num = ""
        current_title = ""
        current_status = ""
        current_prog = ""
        current_entries: list[str] = []

        def _flush():
            if current_num and current_entries:
                rows.append(
                    (
                        int(current_num),
                        current_title,
                        current_status,
                        current_prog,
                        "\n".join(current_entries),
                    )
                )

        for line in text.split("\n"):
            stripped = line.strip()
            if not stripped or stripped.startswith("#"):
                continue
            # 有序列表项："1. title [status, prog]:"
            if stripped[0].isdigit() and ". " in stripped:
                _flush()
                current_entries = []
                parts = stripped.split(". ", 1)
                current_num = parts[0]
                rest = parts[1]
                if " [" in rest and "]:" in rest:
                    current_title = rest.split(" [", 1)[0]
                    bracket = rest.split("[", 1)[1].split("]", 1)[0]
                    if ", " in bracket:
                        current_status, current_prog = bracket.split(", ", 1)
                    else:
                        current_status, current_prog = bracket, ""
                else:
                    current_title = rest.rstrip(":")
            elif line.startswith("    ") and current_num:
                current_entries.append(stripped)

        _flush()
        return rows
